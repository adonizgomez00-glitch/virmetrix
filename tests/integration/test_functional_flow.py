import io
from pathlib import Path

from docx import Document
from docx.shared import Pt

from virmetrix.application.orchestrator import EvaluationOrchestrator
from virmetrix.infrastructure.catalog.loader import load_exercise

EX = "word-format-001"
TARGET = "Informe de resultados"


def _make_correct_bytes():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("El presente ")
    r = p.add_run(TARGET)
    r.bold = True
    r.font.size = Pt(16)
    p.add_run(" resume los hallazgos del ejercicio.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_functional_reference_unformatted(tmp_path):
    orch = EvaluationOrchestrator()
    ref = Path("exercises/word-format-001/reference.docx")
    work = orch.create_work_copy(str(ref), str(tmp_path / "w"))
    ex = load_exercise(EX)
    res = orch.evaluate("docx_portable", work, ex)
    assert res.status == "completed"
    states = {c.criterion_id: c.status.value for c in res.criteria}
    # reference unformatted: C2/C3 not_verifiable or fail; C1 pass; C5 pass
    assert states["C1"] == "pass"
    assert states["C5"] == "pass"


def test_functional_correct_document(tmp_path):
    orch = EvaluationOrchestrator()
    p = tmp_path / "correct.docx"
    p.write_bytes(_make_correct_bytes())
    ex = load_exercise(EX)
    res = orch.evaluate("docx_portable", p, ex)
    assert res.status == "completed"
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states == {"C1": "pass", "C2": "pass", "C3": "pass", "C4": "pass", "C5": "pass"}
    assert res.score_earned == 5


def test_functional_reproducibility(tmp_path):
    p = tmp_path / "correct.docx"
    p.write_bytes(_make_correct_bytes())
    ex = load_exercise(EX)
    orch = EvaluationOrchestrator()
    r1 = orch.evaluate("docx_portable", str(p), ex)
    r2 = orch.evaluate("docx_portable", str(p), ex)
    assert r1.score_earned == r2.score_earned
    assert [(c.criterion_id, c.status.value) for c in r1.criteria] == [(c.criterion_id, c.status.value) for c in r2.criteria]


def test_functional_mode_unavailable_word_com(tmp_path):
    import sys
    if sys.platform == "win32":
        return
    orch = EvaluationOrchestrator()
    avail, reason = orch.available_modes()["word_com"]
    assert avail is False
    assert reason
