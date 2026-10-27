import hashlib
import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from virmetrix.domain.ports import ExerciseMetadata
from virmetrix.infrastructure.catalog.loader import load_exercise
from virmetrix.infrastructure.evaluators.portable import PortableDocxEvaluator

EX = "word-format-001"
TARGET = "Informe de resultados"
BASE_TEXT = "El presente Informe de resultados resume los hallazgos del ejercicio."

evaluator = PortableDocxEvaluator()
exercise = load_exercise(EX)


def _make_doc(runs_spec):
    """runs_spec: list of (text, bold, size_pt_or_None)"""
    doc = Document()
    p = doc.add_paragraph()
    for text, bold, size in runs_spec:
        r = p.add_run(text)
        if bold is not None:
            r.bold = bold
        if size is not None:
            r.font.size = Pt(size)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _write_tmp(tmp_path, data, name="work.docx"):
    p = tmp_path / name
    p.write_bytes(data)
    return str(p)


def test_correct_all_pass(tmp_path):
    data = _make_doc([
        ("El presente ", None, None),
        (TARGET, True, 16),
        (" resume los hallazgos del ejercicio.", None, None),
    ])
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    assert res.status == "completed"
    assert res.score_earned == 5
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states == {"C1": "pass", "C2": "pass", "C3": "pass", "C4": "pass", "C5": "pass"}


def test_missing_bold_fails_c2(tmp_path):
    data = _make_doc([
        ("El presente ", None, None),
        (TARGET, False, 16),
        (" resume los hallazgos del ejercicio.", None, None),
    ])
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C2"] == "fail"
    assert states["C1"] == "pass"
    assert states["C5"] == "pass"


def test_wrong_size_fails_c3(tmp_path):
    data = _make_doc([
        ("El presente ", None, None),
        (TARGET, True, 12),
        (" resume los hallazgos del ejercicio.", None, None),
    ])
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C3"] == "fail"


def test_target_absent_fails_c1(tmp_path):
    data = _make_doc([("Sin la frase objetivo.", None, None)])
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C1"] == "fail"


def test_fragmented_target_passes(tmp_path):
    # split target into multiple runs, all bold+16
    parts = ["Informe", " de ", "resultados"]
    runs = [("El presente ", None, None)]
    for pt in parts:
        runs.append((pt, True, 16))
    runs.append((" resume los hallazgos del ejercicio.", None, None))
    data = _make_doc(runs)
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C1"] == "pass"
    assert states["C2"] == "pass"
    assert states["C3"] == "pass"


def test_text_changed_fails_c5(tmp_path):
    data = _make_doc([
        ("El presente ", None, None),
        (TARGET, True, 16),
        (" resume los hallazgos modificados.", None, None),
    ])
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C5"] == "fail"


def test_corrupt_file_returns_error(tmp_path):
    path = _write_tmp(tmp_path, b"not a zip", name="bad.docx")
    res = evaluator.evaluate(path, exercise)
    assert res.status == "error"
    assert res.error_message


def test_neighbor_format_leak_fails_c4(tmp_path):
    # neighbor also bold+16 -> C4 fail
    data = _make_doc([
        ("El presente ", True, 16),
        (TARGET, True, 16),
        (" resume los hallazgos del ejercicio.", None, None),
    ])
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C4"] == "fail"


def test_multiple_occurrences_fail_c1(tmp_path):
    data = _make_doc([
        ("Informe de resultados ", None, None),
        (TARGET, True, 16),
        (" resume los hallazgos del ejercicio.", None, None),
    ])
    path = _write_tmp(tmp_path, data)
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C1"] == "fail"


def test_inherited_not_verifiable(tmp_path):
    # target run has no direct bold/sz (inherit) -> not_verifiable for C2/C3
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run(TARGET)
    # no direct formatting
    buf = io.BytesIO()
    # add surrounding text
    p2 = doc.paragraphs[0]
    # rebuild: insert before/after by creating new doc
    doc2 = Document()
    pp = doc2.add_paragraph()
    pp.add_run("El presente ")
    pp.add_run(TARGET)  # no formatting
    pp.add_run(" resume los hallazgos del ejercicio.")
    buf = io.BytesIO()
    doc2.save(buf)
    path = _write_tmp(tmp_path, buf.getvalue())
    res = evaluator.evaluate(path, exercise)
    states = {c.criterion_id: c.status.value for c in res.criteria}
    assert states["C2"] == "not_verifiable"
    assert states["C3"] == "not_verifiable"
