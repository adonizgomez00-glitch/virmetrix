import io
import tempfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from virmetrix.infrastructure.catalog.loader import list_exercises, load_exercise
from virmetrix.infrastructure.evaluators.portable import PortableDocxEvaluator

ev = PortableDocxEvaluator()


def _correct_bytes(ex):
    prop = None
    size = None
    for s in ex.criteria_specs:
        if s["type"] == "run_property" and s["property"] in ("bold", "italic"):
            prop = s["property"]
        if s["type"] == "run_property" and s["property"] == "size_pt":
            size = s["expected"]
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("El presente ")
    r = p.add_run("Informe de resultados")
    if prop == "bold":
        r.bold = True
    if prop == "italic":
        r.italic = True
    r.font.size = Pt(size)
    p.add_run(" resume los hallazgos del ejercicio.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("exinfo", [e for e in list_exercises() if e["exercise_id"].startswith("word-format")], ids=lambda e: e["exercise_id"])
def test_catalog_exercise_correct_passes(exinfo, tmp_path):
    ex = load_exercise(exinfo["exercise_id"])
    data = _correct_bytes(ex)
    path = tmp_path / f"{ex.exercise_id}.docx"
    path.write_bytes(data)
    res = ev.evaluate(str(path), ex)
    assert res.status == "completed"
    assert res.score_earned == res.score_possible
    assert all(c.status.value == "pass" for c in res.criteria)


def test_table_001_correct(tmp_path):
    from docx import Document
    ex = load_exercise("table-001")
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Producto"
    table.cell(0, 1).text = "Cantidad"
    table.cell(1, 0).text = "Teclado"
    table.cell(1, 1).text = "5"
    buf = io.BytesIO()
    doc.save(buf)
    path = tmp_path / "table.docx"
    path.write_bytes(buf.getvalue())
    res = ev.evaluate(str(path), ex)
    assert res.score_earned == res.score_possible


def test_list_001_correct(tmp_path):
    from docx import Document
    ex = load_exercise("list-001")
    doc = Document()
    doc.add_paragraph("Lista de compras")
    for t in ["Leche", "Pan", "Huevos"]:
        doc.add_paragraph(t, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    path = tmp_path / "list.docx"
    path.write_bytes(buf.getvalue())
    res = ev.evaluate(str(path), ex)
    assert res.score_earned == res.score_possible


def test_list_exercises_not_empty():
    assert len(list_exercises()) >= 4