import io
import zipfile
from pathlib import Path

from virmetrix.domain.ports import ExerciseMetadata
from virmetrix.infrastructure.catalog.loader import load_exercise
from virmetrix.infrastructure.evaluators.portable import PortableDocxEvaluator, MAX_FILE_SIZE

exercise = load_exercise("word-format-001")
ev = PortableDocxEvaluator()


def _write(tmp_path, data, name="x.docx"):
    p = tmp_path / name
    p.write_bytes(data)
    return str(p)


def test_traversal_path_returns_error(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../evil.xml", "<x/>")
        # also include document.xml to pass initial? traversal check happens before read
    path = _write(tmp_path, buf.getvalue())
    res = ev.evaluate(path, exercise)
    assert res.status == "error"


def test_traversal_backslash_returns_error(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("..\\evil.xml", "<x/>")
    path = _write(tmp_path, buf.getvalue())
    res = ev.evaluate(path, exercise)
    assert res.status == "error"


def test_invalid_zip_returns_error(tmp_path):
    path = _write(tmp_path, b"not a zip file at all")
    res = ev.evaluate(path, exercise)
    assert res.status == "error"


def test_missing_document_xml_returns_error(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("other.txt", "hi")
    path = _write(tmp_path, buf.getvalue())
    res = ev.evaluate(path, exercise)
    assert res.status == "error"


def test_empty_file_returns_error(tmp_path):
    path = _write(tmp_path, b"")
    res = ev.evaluate(path, exercise)
    assert res.status == "error"


def test_nonexistent_file_returns_error(tmp_path):
    res = ev.evaluate(str(tmp_path / "nope.docx"), exercise)
    assert res.status == "error"


def test_unicode_filename(tmp_path):
    p = tmp_path / "trabajo (copy) áéíóú.docx"
    # create a valid docx
    from docx import Document
    doc = Document()
    doc.add_paragraph("hola")
    doc.save(str(p))
    res = ev.evaluate(str(p), exercise)
    # should not crash; status completed or error controlled
    assert res.status in ("completed", "error")


def test_not_verifiable_not_counted(tmp_path):
    # inherited (no direct formatting) -> C2/C3 not_verifiable; score should not count them
    from docx import Document
    doc = Document()
    pp = doc.add_paragraph()
    pp.add_run("El presente ")
    pp.add_run("Informe de resultados")
    pp.add_run(" resume los hallazgos del ejercicio.")
    p = tmp_path / "inh.docx"
    doc.save(str(p))
    res = ev.evaluate(str(p), exercise)
    assert res.status == "completed"
    # not_verifiable should not be counted as pass
    for c in res.criteria:
        if c.status.value == "not_verifiable":
            assert c.criterion_id in ("C2", "C3")
    assert res.score_earned <= 3
